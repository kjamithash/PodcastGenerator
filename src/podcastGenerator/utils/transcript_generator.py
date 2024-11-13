from langchain import PromptTemplate, LLMChain
from langchain_anthropic import ChatAnthropic
from docx import Document
import logging
import os
from pathlib import Path
import time
from tenacity import (
    retry,
    stop_after_attempt,
    wait_exponential,
    retry_if_exception_type
)

def setup_logging():
    """Configure logging for the transcript generator."""
    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
    )

def clean_transcript(transcript: str, model_name: str) -> str:
    """
    Clean the transcript by removing unwanted lines and ensuring proper outro.
    
    Args:
        transcript (str): Raw transcript from the LLM
        model_name (str): Name of the mental model
        
    Returns:
        str: Cleaned transcript
    """
    # Split into lines for processing
    lines = transcript.split('\n')
    
    # Remove the first line if it contains helper text
    if any(phrase in lines[0].lower() for phrase in [
        "let me help",
        "i'll create",
        "here's a transcript",
        "i'll write",
        "let me write"
    ]):
        lines = lines[1:]
    
    # Replace outro placeholder with standard outro
    standard_outro = (
        "For more mental models, please visit mentalmodelsdaily.com or "
        "find us on X or Instagram. Our Podcast music was provided by "
        "thePodcasthost.com & Alitu: The Podcast Maker. Find your own free "
        "podcast music over at thePodcasthost.com/freemusic."
    )
    
    # Remove any existing outro placeholders or incomplete outros
    while lines and any(phrase in lines[-1].lower() for phrase in [
        "[outro",
        "outro:",
        "for more mental models",
        "mentalmodelsdaily.com"
    ]):
        lines.pop()
    
    # Add the standard outro
    lines.append(standard_outro)
    
    # Join lines back together
    cleaned_transcript = '\n'.join(lines).strip()
    
    return cleaned_transcript

def get_prompt_template() -> PromptTemplate:
    """Create the prompt template for transcript generation."""
    template = """You are a podcast host for Mental Models Daily, where you explain one mental model each day. 
Your task is to create a transcript for a podcast episode about {model_name}.

Follow this exact structure and style from the Super Forecasters example:

1. Start with a warm welcome and introduce the mental model in an engaging way that hooks the listener
2. Explain the core concept clearly and simply, using a relatable analogy ("It's like...")
3. Share a compelling historical example that illustrates the model
4. Provide a modern business example showing practical application
5. Give three specific ways to apply this in daily life, each with a relatable analogy
6. Conclude by summarizing the key insights and value of the model
7. End with the standard sign-off

Use a conversational, engaging tone throughout. Include natural transitions between sections and "It's like..." analogies to make concepts more relatable.

Each section should be separated by: <break time="{break_duration}" />

Here's your transcript for {model_name}:"""

    return PromptTemplate(
        input_variables=["model_name", "break_duration"],
        template=template
    )

def save_transcript(transcript: str, output_path: str, model_name: str):
    """Save the transcript to a Word document."""
    doc = Document()
    doc.add_paragraph(transcript)
    
    os.makedirs(output_path, exist_ok=True)
    
    model_name_lower = model_name.lower().replace(" ", "-")
    file_path = os.path.join(output_path, f"{model_name_lower}_transcript.docx")
    doc.save(file_path)
    logging.info(f"Transcript saved to {file_path}")

@retry(
    retry=retry_if_exception_type((Exception)),
    wait=wait_exponential(multiplier=1, min=4, max=60),
    stop=stop_after_attempt(5),
    before_sleep=lambda retry_state: logging.info(
        f"Retrying after attempt {retry_state.attempt_number} in {retry_state.next_action.sleep} seconds..."
    )
)
def generate_single_transcript(llm_chain, model_name: str, break_duration: str) -> str:
    """Generate a single transcript with retry logic."""
    try:
        transcript = llm_chain.run({
            "model_name": model_name,
            "break_duration": break_duration
        })
        
        # Ensure sign-off is included
        if "For more mental models, please visit mentalmodelsdaily.com" not in transcript:
            sign_off = ("\nFor more mental models, please visit mentalmodelsdaily.com or "
                      "find us on X or Instagram. Our Podcast music was provided by "
                      "thePodcasthost.com & Alitu: The Podcast Maker. Find your own free "
                      "podcast music over at thePodcasthost.com/freemusic.")
            transcript += sign_off
            
        return transcript
        
    except Exception as e:
        logging.error(f"Error generating transcript: {str(e)}")
        raise

def generate_transcripts(mental_models, output_path):
    """
    Generate transcripts for the given mental models with improved error handling.
    
    Args:
        mental_models (list): List of mental model names to generate transcripts for
        output_path (str): Path where transcript files should be saved
        
    Returns:
        dict: Dictionary mapping model names to their generated transcripts
    """
    setup_logging()
    
    # Initialize the LLM with increased timeout and retries
    llm = ChatAnthropic(
        model="claude-3-5-sonnet-latest",
        temperature=0.7,
        max_retries=3,
        timeout=300  # 5 minute timeout
    )
    
    # Create the chain
    llm_chain = LLMChain(
        prompt=get_prompt_template(),
        llm=llm
    )
    
    transcripts = {}
    break_duration = "1.3s"
    
    for model in mental_models:
        try:
            logging.info(f"Starting transcript generation for {model}")
            
            # Generate with retry logic
            transcript = generate_single_transcript(llm_chain, model, break_duration)

            # Clean the transcript
            cleaned_transcript = clean_transcript(transcript, model)
            
            # Save successful transcript
            save_transcript(cleaned_transcript, output_path, model)
            transcripts[model] = cleaned_transcript
            
            logging.info(f"Successfully generated and saved transcript for {model}")
            
            # Add a delay between requests to avoid rate limiting
            if model != mental_models[-1]:  # Don't delay after the last model
                time.sleep(2)
                
        except Exception as e:
            logging.error(f"Failed to generate transcript for {model} after all retries: {str(e)}")
            continue
    
    return transcripts